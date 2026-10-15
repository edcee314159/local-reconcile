from __future__ import annotations
import csv, hashlib, io, re, zipfile
from datetime import date, datetime
from pathlib import Path
from concurrent.futures import CancelledError
from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from .domain import (RawTable, RawRow, RawCell, LocalReconcileError, check_cancel,
                     quiet_progress, never_cancel)
from . import input_limits as limits

def _bytes(path):
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".csv"):
        raise LocalReconcileError("file_type", "仅支持 .xlsx 和 UTF-8 CSV，请先另存格式")
    try:
        if path.stat().st_size > limits.MAX_FILE_BYTES:
            raise LocalReconcileError("file_limit", "文件超过 25 MiB，请拆分文件")
        content = path.read_bytes()
    except OSError:
        raise LocalReconcileError("file_read", "无法读取文件，请检查路径与权限") from None
    if not content:
        raise LocalReconcileError("empty_file", "文件为空")
    return path, content

def _zip_check(content, cancelled):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if len(infos) > limits.MAX_ZIP_MEMBERS or sum(i.file_size for i in infos) > limits.MAX_UNPACKED_BYTES:
            raise LocalReconcileError("zip_limit", "工作簿解压数据量超过限制")
        for info in infos:
            check_cancel(cancelled)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                with archive.open(info) as stream:
                    for _, element in iterparse(stream, events=("end",)):
                        if element.tag.endswith("}c"):
                            ref = element.get("r", "")
                            m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
                            if not m or column_index_from_string(m[1]) > limits.MAX_COLUMNS or int(m[2]) > limits.MAX_ROWS + 1000:
                                raise LocalReconcileError("sheet_limit", "工作表行列范围超过限制，请整理后导入")
                        element.clear()

def inspect_file(path, role):
    path, content = _bytes(path)
    if path.suffix.lower() == ".csv":
        return ()
    try:
        _zip_check(content, never_cancel)
        book = load_workbook(io.BytesIO(content), read_only=True, keep_links=False)
        try:
            return tuple(book.sheetnames)
        finally:
            book.close()
    except LocalReconcileError:
        raise
    except Exception:
        raise LocalReconcileError("workbook", "无法打开工作簿，请检查文件是否损坏或加密") from None

def _cell(value, cached=None):
    if value is None:
        kind = "empty"
    elif isinstance(value, (date, datetime)):
        kind = "date"
    elif isinstance(value, (int, float, bool)):
        kind = "number"
    else:
        kind = "text"
    if isinstance(value, str) and len(value) > limits.MAX_CELL_CHARS:
        raise LocalReconcileError("cell_limit", "单元格文本超过 32767 字符")
    return RawCell(value, kind, cached)

def read_table(path, role, sheet=None, header_row=1, progress=quiet_progress, cancelled=never_cancel):
    check_cancel(cancelled)
    if role not in ("bank", "ledger") or type(header_row) is not int or not 1 <= header_row <= 1000:
        raise LocalReconcileError("mapping", "文件角色或表头行无效")
    path, content = _bytes(path)
    fingerprint = hashlib.sha256(content).hexdigest()
    rows = []
    headers = None
    epoch = None
    def consume(line, cells):
        nonlocal headers
        check_cancel(cancelled)
        if len(cells) > limits.MAX_COLUMNS:
            raise LocalReconcileError("column_limit", "列数超过 200，请整理工作表")
        if line < header_row:
            return
        if headers is None:
            if line != header_row:
                raise LocalReconcileError("header", "所选表头行位于多行文本内部")
            headers = tuple(str(c.value or "").strip() for c in cells)
            return
        if len(rows) >= limits.MAX_ROWS:
            raise LocalReconcileError("row_limit", "数据超过 100000 行，请分批处理")
        row_id = hashlib.sha256(f"{role}|{fingerprint}|{sheet}|{line}".encode()).hexdigest()[:24]
        rows.append(RawRow(row_id, line, tuple(cells)))
        if len(rows) % 250 == 0:
            progress("读取文件", len(rows), 0)
    try:
        if path.suffix.lower() == ".csv":
            csv.field_size_limit(limits.MAX_CELL_CHARS)
            reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
            while True:
                start = reader.line_num + 1
                try:
                    values = next(reader)
                except StopIteration:
                    break
                consume(start, [_cell(v) for v in values])
        else:
            _zip_check(content, cancelled)
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False, keep_vba=False)
            cache = None
            try:
                cache = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False, keep_vba=False)
                sheet = sheet or book.sheetnames[0]
                if sheet not in book.sheetnames:
                    raise LocalReconcileError("sheet", "所选工作表不存在")
                ws, cached = book[sheet], cache[sheet]
                ws.reset_dimensions()
                cached.reset_dimensions()
                epoch = book.epoch.date()
                for line, (values, cached_values) in enumerate(zip(ws.iter_rows(), cached.iter_rows()), 1):
                    cells = []
                    for index, c in enumerate(values):
                        cv = cached_values[index].value if index < len(cached_values) else None
                        if c.data_type == "f":
                            cells.append(RawCell(c.value, "formula", cv))
                        else:
                            cells.append(_cell(c.value))
                    consume(line, cells)
            finally:
                book.close()
                if cache is not None:
                    cache.close()
    except (LocalReconcileError, CancelledError):
        raise
    except Exception:
        raise LocalReconcileError("file_parse", "读取失败，请检查格式、编码、文件完整性与加密状态") from None
    if not headers or not any(headers) or not rows:
        raise LocalReconcileError("empty_data", "所选表头下没有可读取的数据")
    width = max(len(headers), max(len(row.cells) for row in rows))
    headers = headers + tuple(f"未命名列 {i+1}" for i in range(len(headers), width))
    progress("读取完成", len(rows), len(rows))
    return RawTable(role, path.name, fingerprint, sheet, header_row, epoch, headers, tuple(rows))
