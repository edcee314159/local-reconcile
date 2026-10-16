"""Traceable Excel output. Imported strings never become executable formulas."""
from pathlib import Path
from decimal import Decimal
from datetime import date,datetime
from concurrent.futures import CancelledError
import os,tempfile
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
from .domain import LocalReconcileError,check_cancel,never_cancel
from .amounts import format_amount
from .review import classify_rows
from .project_store import protect_input_path

def _money(value):
    return Decimal(value)/100 if abs(value)<10**15 else format_amount(value)

_BODY_FONT=Font(name="Microsoft YaHei",size=10)
_BODY_ALIGNMENT=Alignment(vertical="top")

def _append(sheet,values):
    cells=[]
    for value in values:
        if isinstance(value,(date,datetime)):value=value.isoformat()
        cell=WriteOnlyCell(sheet,value=value)
        if isinstance(value,str):cell.data_type="s";cell.number_format="@"
        elif isinstance(value,Decimal):cell.number_format='#,##0.00;[Red]-#,##0.00'
        cell.font=_BODY_FONT;cell.alignment=_BODY_ALIGNMENT;cells.append(cell)
    sheet.append(cells)

def export_report(snapshot,path,cancelled=never_cancel):
    path=Path(path);temporary=None;workbook=None
    check_cancel(cancelled);protect_input_path(snapshot,path)
    if path.suffix.lower()!=".xlsx":
        raise LocalReconcileError("export_type","报告请保存为 .xlsx 文件")
    try:
        states=classify_rows(snapshot)
        bank={t.id:t for t in snapshot.bank.transactions};ledger={t.id:t for t in snapshot.ledger.transactions}
        br={r.id:r for r in snapshot.bank.raw.rows};lr={r.id:r for r in snapshot.ledger.raw.rows}
        workbook=Workbook();workbook.remove(workbook.active)
        names=("核对总览","已确认关系","待确认候选","歧义与未匹配","银行原始明细","日记账原始明细","输入检查","参数与操作记录")
        sheets={n:workbook.create_sheet(n) for n in names}
        def add(name,row):_append(sheets[name],row)
        add("核对总览",["项目","银行流水","银行日记账"])
        add("核对总览",["交易笔数",len(bank),len(ledger)])
        for label,status in (("已确认笔数","confirmed"),("有候选待复核","pending"),("无候选未匹配","unmatched")):
            add("核对总览",[label,*[sum(states[t.id]==status for t in d.transactions) for d in (snapshot.bank,snapshot.ledger)]])
        for label,calc in (("收入合计",lambda t:max(t.amount_minor,0)),("支出合计",lambda t:max(-t.amount_minor,0)),("净额",lambda t:t.amount_minor)):
            add("核对总览",[label,*[_money(sum(calc(t) for t in d.transactions)) for d in (snapshot.bank,snapshot.ledger)]])
        add("核对总览",["已确认净额",*[_money(sum(t.amount_minor for t in d.transactions if states[t.id]=="confirmed")) for d in (snapshot.bank,snapshot.ledger)]])
        add("核对总览",["口径","人民币；收入正、支出负；候选不计入确认合计","结果需人工复核，不自动过账"])
        add("核对总览",["原始记录行数",len(br),len(lr)])
        add("核对总览",["日期容差（天）",snapshot.config.date_tolerance_days,"规则版本 "+snapshot.result.rule_version])
        add("已确认关系",["关系编号","银行文件","工作表","源行","银行记录编号","日记账文件","工作表","源行","日记账记录编号","银行日期","日记账日期","银行金额（元）","日记账金额（元）","确认依据"])
        for p in snapshot.review.pairs:
            b,l=bank[p.bank_id],ledger[p.ledger_id]
            add("已确认关系",[p.id,snapshot.bank.raw.filename,snapshot.bank.raw.sheet or "CSV",br[b.id].source_line,b.id,snapshot.ledger.raw.filename,snapshot.ledger.raw.sheet or "CSV",lr[l.id].source_line,l.id,b.day,l.day,_money(b.amount_minor),_money(l.amount_minor),p.basis])
        add("待确认候选",["候选编号","银行记录编号","日记账记录编号","银行源行","日记账源行","银行日期","日记账日期","金额（元）","规则","冲突","当前状态"])
        rejected=set(snapshot.review.rejected);confirmed={(p.bank_id,p.ledger_id) for p in snapshot.review.pairs}
        for n,c in enumerate(snapshot.result.candidates):
            if n%250==0:check_cancel(cancelled)
            status="已确认" if (c.bank_id,c.ledger_id) in confirmed else "已拒绝" if (c.bank_id,c.ledger_id) in rejected else "记录已被占用" if states[c.bank_id]=="confirmed" or states[c.ledger_id]=="confirmed" else "待复核"
            b,l=bank[c.bank_id],ledger[c.ledger_id]
            add("待确认候选",[c.id,b.id,l.id,br[b.id].source_line,lr[l.id].source_line,b.day,l.day,_money(b.amount_minor),c.rule,",".join(c.conflicts),status])
        add("歧义与未匹配",["类型","银行记录编号","日记账记录编号","银行源行","日记账源行","金额或差额（元）","说明","分组编号"])
        for d in snapshot.result.differences:
            add("歧义与未匹配",["金额差异",d.bank_id,d.ledger_id,br[d.bank_id].source_line,lr[d.ledger_id].source_line,_money(d.delta_minor),d.reason,""])
        for ds,transactions,rawrows in ((snapshot.bank,bank,br),(snapshot.ledger,ledger,lr)):
            isbank=ds.raw.role=="bank"
            for rid,t in transactions.items():
                if states[rid]=="unmatched":
                    add("歧义与未匹配",["未匹配",rid if isbank else "",rid if not isbank else "",rawrows[rid].source_line if isbank else "",rawrows[rid].source_line if not isbank else "",_money(t.amount_minor),ds.raw.role+" 未匹配",""])
        for g in snapshot.result.groups:
            for rid in g.bank_ids+g.ledger_ids:
                isbank=rid in bank;t=bank[rid] if isbank else ledger[rid];r=br[rid] if isbank else lr[rid]
                add("歧义与未匹配",["范围受限" if g.truncated else "歧义",rid if isbank else "",rid if not isbank else "",r.source_line if isbank else "",r.source_line if not isbank else "",_money(t.amount_minor),g.reason,g.id])
        for ds,name in ((snapshot.bank,"银行原始明细"),(snapshot.ledger,"日记账原始明细")):
            add(name,["记录编号","源文件","工作表","源行",*ds.raw.headers])
            for n,r in enumerate(ds.raw.rows):
                if n%250==0:check_cancel(cancelled)
                add(name,[r.id,ds.raw.filename,ds.raw.sheet or "CSV",r.source_line,*[c.value for c in r.cells]])
        add("输入检查",["代码","级别","说明","来源记录编号"])
        for i in snapshot.result.issues:
            for rid in i.row_ids or ("",):add("输入检查",[i.code,i.severity,i.message,rid])
        add("参数与操作记录",["类型","键或序号","值或操作","银行记录编号","日记账记录编号","备注"])
        for key,value in (("规则版本",snapshot.result.rule_version),("日期容差",snapshot.config.date_tolerance_days),("候选展开上限",snapshot.config.max_pairs_per_group),("银行文件 SHA256",snapshot.bank.raw.sha256),("日记账 SHA256",snapshot.ledger.raw.sha256)):
            add("参数与操作记录",["参数",key,value,"","",""])
        for e in snapshot.review.events:add("参数与操作记录",["人工操作",e.sequence,e.kind,e.bank_id,e.ledger_id,e.note])
        for sheet in workbook:
            sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.showGridLines=False
            sheet.row_dimensions[1].height=28
            for c in sheet[1]:
                c.fill=PatternFill("solid",fgColor="193B50");c.font=Font(name="Microsoft YaHei",size=10,bold=True,color="FFFFFF")
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width=min(45,max(15,max(len(str(c.value or "")) for c in col[:100])+2))
        check_cancel(cancelled)
        with tempfile.NamedTemporaryFile(dir=path.parent,suffix=".tmp",delete=False) as f:temporary=Path(f.name)
        workbook.save(temporary);check_cancel(cancelled);os.replace(temporary,path)
    except (LocalReconcileError,CancelledError):raise
    except Exception:raise LocalReconcileError("export","导出失败，请检查文件占用、目录权限与剩余空间") from None
    finally:
        if workbook:workbook.close()
        if temporary and temporary.exists():temporary.unlink(missing_ok=True)
