from dataclasses import asdict
from hashlib import sha256
from pathlib import Path
import json,os,tempfile
from .models import Account,Batch,Settings,Event,Workspace
from .workspace import new_workspace,add_account,add_batch,record_event,replay,period_end,active_records,fail
from .operations import open_items,outstanding
from .settlements import advance_rows, allocation_net
from .credits import credit_rows
from .duplicates import duplicate_groups
from .counterparty import counterparty_report
from .reporting import bank_statement,aging_report,AGE_BUCKETS,TIMING_CATEGORIES
from ..project_store import _encode,_decode,validate_snapshot
from ..amounts import format_amount
from ..domain import LocalReconcileError,check_cancel,never_cancel

MAX_BYTES=100*1024*1024

def _canonical(value):
    return json.dumps(value,ensure_ascii=False,sort_keys=True,allow_nan=False,separators=(",",":")).encode("utf-8")

def _protect(w,path):
    path=Path(path)
    if path.name.casefold() in {b.raw.filename.casefold() for b in w.batches}:fail("输出不得覆盖或沿用原始输入文件名")
    if path.is_file() and path.stat().st_size<=25*1024*1024:
        if sha256(path.read_bytes()).hexdigest() in {b.raw.sha256 for b in w.batches}:fail("输出目标属于原始输入，请另存")

def _write_atomic(path,data,cancelled=never_cancel):
    path=Path(path);temporary=None
    try:
        check_cancel(cancelled)
        with tempfile.NamedTemporaryFile(dir=path.parent,suffix=".tmp",delete=False) as f:
            temporary=Path(f.name);f.write(data);f.flush();os.fsync(f.fileno())
        check_cancel(cancelled);os.replace(temporary,path)
    finally:
        if temporary:temporary.unlink(missing_ok=True)

def save_workspace(w,path,cancelled=never_cancel):
    _protect(w,path);replay(w)
    batches=[]
    for b in w.batches:
        body={"id":b.id,"role":b.role,"account":b.account,"raw":_encode(b.raw),"mapping":_encode(b.mapping),"due_col":b.due_col}
        if b.statement_period is not None:body["statement_period"]=b.statement_period
        if b.extra_cols:body["extra_cols"]=b.extra_cols
        batches.append({**body,"payload_sha256":sha256(_canonical(body)).hexdigest()})
    data={"version":2,"start_period":w.start_period,"accounts":[asdict(a) for a in w.accounts],
          "settings":asdict(w.settings),"batches":batches,"events":[asdict(e) for e in w.events]}
    try:
        content=_canonical(data)
        if len(content)>MAX_BYTES:fail("业务工作区超过 100 MiB")
        _write_atomic(path,content,cancelled)
    except (OSError,TypeError,ValueError) as exc:
        if isinstance(exc,LocalReconcileError):raise
        fail("业务工作区保存失败，请检查目录、空间与文件占用")

def load_workspace(path):
    try:
        path=Path(path)
        if path.stat().st_size>MAX_BYTES:fail("业务工作区超过 100 MiB")
        d=json.loads(path.read_text(encoding="utf-8"),parse_constant=lambda _:fail("无效数值"))
        if type(d) is not dict or set(d)!={"version","start_period","accounts","settings","batches","events"} or type(d["version"]) is not int or d["version"]!=2:fail("业务工作区版本或字段无效")
        w=new_workspace(d["start_period"])
        for a in d["accounts"]:
            if type(a) is not dict or set(a)!={"code","label"}:fail("账户字段无效")
            w=add_account(w,**a)
        batches=[]
        for b in d["batches"]:
            if type(b) is not dict or set(b)-{"statement_period","extra_cols"}!={"id","role","account","raw","mapping","due_col","payload_sha256"}:fail("来源字段无效")
            body={k:v for k,v in b.items() if k!="payload_sha256"}
            if sha256(_canonical(body)).hexdigest()!=b["payload_sha256"]:fail("来源内容校验失败")
            batches.append(Batch(b["id"],b["role"],b["account"],_decode(b["raw"]),_decode(b["mapping"]),b["due_col"],b.get("statement_period"),b.get("extra_cols",{})))
        if type(d["events"]) is not list or len(d["events"])>200000:fail("操作记录超过限制")
        events=[]
        for e in d["events"]:
            if type(e) is not dict or set(e)!={"seq","kind","data"}:fail("操作记录字段无效")
            events.append(Event(**e))
        w=Workspace(w.start_period,w.accounts,tuple(batches),tuple(events),Settings(**d["settings"]))
        w.records;w.state
        return w
    except LocalReconcileError:raise
    except Exception:fail("工作区无法读取，文件损坏或结构无效；原文件未修改")

def import_legacy(snapshot,period):
    validate_snapshot(snapshot)
    code=snapshot.bank.mapping.declared_account
    w=add_account(new_workspace(period),code,code)
    w=add_batch(w,"bank",code,snapshot.bank.raw,snapshot.bank.mapping)
    w=add_batch(w,"ledger",code,snapshot.ledger.raw,snapshot.ledger.mapping)
    lookup={r.source_id:r.id for r in w.records.values()}
    for e in snapshot.review.events:
        left=[lookup[e.bank_id]];right=[lookup[e.ledger_id]]
        note=e.note or "旧项目人工复核记录"
        if e.kind in ("confirm","reject"):
            w=record_event(w,"confirm_match" if e.kind=="confirm" else "reject_match",left=left,right=right,note=note,legacy_sequence=e.sequence)
        else:
            target=next((k for k,p in w.state.matches.items() if p["left"]==tuple(left) and p["right"]==tuple(right)),None)
            if target:w=record_event(w,"undo_match",target=target,note=note,legacy_sequence=e.sequence)
            else:w=record_event(w,"undo_reject",left=left,right=right,note=note,legacy_sequence=e.sequence)
    return w

def export_workspace(w,path,cancelled=never_cancel):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from ..export_xlsx import _append
    path=Path(path);_protect(w,path)
    if path.suffix.lower()!=".xlsx":fail("报告请保存为 .xlsx")
    check_cancel(cancelled);s=replay(w)
    wb=Workbook();wb.remove(wb.active);sheets={}
    headers={
      "核算总览":["项目","数值","说明"],"已确认对账":["关系编号","账户","银行笔数","账面笔数","确认净额（元）","依据","期间"],
      "待确认候选":["候选编号","银行行编号","账面行编号","合计（元）","依据","冲突","状态"],
      "搜索范围":["范围编号","说明","来源记录编号","记录信息"],
      "对账组成明细":["关系编号","来源","记录编号","日期","金额（元）","交易对方","共同编号","原文件","源行"],
      "差额说明":["说明编号","银行记录","账面记录","差额（元）","类别","结算依据","状态","期间"],
      "未达明细":["记录编号","来源","账户","日期","金额（元）","对方","编号","账龄天数","已有差额说明"],
      "期间快照":["期间","类型","记录或关系编号","余额（元）"],
      "往来核销":["核销编号","银行记录","往来单据记录","原核销金额（元）","依据","期间","退款后净核销（元）","预收预付登记"],
      "往来余额":["记录编号","类型","往来单位","单据编号","单据日期","单据金额（元）","资金净核销（元）","红字抵减（元）","未结清（元）","账龄天数","到期日","逾期天数"],
      "内部转账":["关系编号","转出记录","转出账户","转入记录","转入账户","转入金额（元）","手续费（元）","依据","期间"],
      "退款冲销":["关系编号","原银行记录","退款流水","金额（元）","类别","关联核销","关联预收预付","依据","期间"],
      "预收预付余额":["登记编号","类型","银行记录","往来单位","原登记金额（元）","未核销余额（元）","依据","期间"],
      "银行余额调节表":["期间","账户","项目","金额（元）","状态或依据"],
      "余额调节明细":["期间","账户","记录编号","来源","日期","原金额（元）","未达类别","已复核","依据","原文件","源行"],
      "账龄汇总":["期间","类型","往来单位","未结清（元）",*AGE_BUCKETS,"逾期金额（元）","未到期金额（元）","缺到期日金额（元）","单据数"],
      "账龄明细":["期间","类型","往来单位","单据编号","记录编号","单据日期","到期日","资金净核销（元）","红字抵减（元）","未结清（元）","账龄天数","账龄区间","逾期天数","状态","原文件","源行"],
      "原始记录":["批次","业务类型","账户","原文件","文件SHA256","工作表","源行","原始行编号","原始字段与内容"],
      "红字抵减":["操作编号","红字记录","原单据记录","抵减金额（元）","类别","依据","期间"],
      "红字可用余额":["记录编号","类型","往来单位","红字编号","原金额（元）","已抵减（元）","可用余额（元）"],
      "疑似重复交易":["期间","组编号","状态","判定条件","记录编号","来源类别","账户","日期","金额（元）","往来单位","编号","原文件","源行","复核依据"],
      "双方往来对账汇总":["期间","对账单批次","往来单位","方向","截止日","我方未结（元）","对方未结（元）","差额（元）","已确认行数","已说明差异行数","待复核行数","对账单依据"],
      "双方往来对账明细":["期间","对账单批次","单据编号","我方记录","对方记录","我方未结（元）","对方未结（元）","差额（元）","状态","复核依据"],
      "操作记录":["序号","操作","内容"]}
    from .advanced_export import HEADERS,export_advanced
    headers.update(HEADERS)
    for name,head in headers.items():sheets[name]=wb.create_sheet(name);_append(sheets[name],head)
    def put(name,row):_append(sheets[name],row)
    batches={b.id:b for b in w.batches};source_rows={b.id:{r.id:r for r in b.raw.rows} for b in w.batches}
    put("核算总览",["当前期间",s.period,"人民币；候选不计入确认金额；本报告不自动过账"])
    put("核算总览",["已确认对账关系",len(s.matches),"组合按关系统计，不做两侧明细笛卡尔积"])
    put("核算总览",["已确认净额（元）",format_amount(sum(p["amount_minor"] for p in s.matches.values())),"每组银行侧只计一次"])
    put("核算总览",["内部转账关系",len(s.transfers),"账户间调拨，不抵消原始银行收付记录"])
    put("核算总览",["原始行数",sum(len(b.raw.rows) for b in w.batches),"保留空白及错误提示的来源信息"])
    put("核算总览",["组合参数",json.dumps(asdict(w.settings),ensure_ascii=False),"规则版本 2；搜索限制不表示已穷尽匹配"])
    for account in w.accounts:
        for role in ("bank","ledger"):
            rows=[r for r in active_records(w,s).values() if r.account==account.code and r.role==role]
            put("核算总览",[account.label+" / "+role+" 净额",format_amount(sum(r.amount_minor for r in rows)),account.code])
    for rid,p in s.matches.items():
        check_cancel(cancelled)
        put("已确认对账",[rid,w.records[p["left"][0]].account,len(p["left"]),len(p["right"]),format_amount(p["amount_minor"]),p["note"] or p["basis"],p["period"]])
        for key in ("left","right"):
            for row_id in p[key]:
                r=w.records[row_id];b=batches[r.batch_id];raw=source_rows[b.id][r.source_id]
                put("对账组成明细",[rid,r.role,r.id,r.day,format_amount(r.amount_minor),r.party,r.reference,b.raw.filename,raw.source_line])
    from .matching import suggest_matches
    analysis=suggest_matches(w,cancelled=cancelled)
    confirmed={(p["left"],p["right"]) for p in s.matches.values()}
    occupied={rid for p in s.matches.values() for rid in p["left"]+p["right"]}
    for c in analysis.candidates:
        key=c.left,c.right
        status="已确认" if key in confirmed else "已拒绝" if key in s.rejected else "记录已占用" if occupied.intersection(c.left+c.right) else "待复核"
        put("待确认候选",[c.id,", ".join(c.left),", ".join(c.right),format_amount(c.amount_minor),c.basis,"; ".join(c.conflicts),status])
    for n,scope in enumerate(analysis.ranges,1):
        for rid in scope.rows:
            r=w.records[rid];put("搜索范围",[n,scope.reason,rid,f"{r.account} / {r.day} / {format_amount(r.amount_minor)} / {r.party} / {r.reference}"])
    for rid,p in s.explanations.items():put("差额说明",[rid,p["bank"],p["ledger"],format_amount(p["delta_minor"]),p["category"],p["evidence"],"待查" if p["category"]=="待查" else "差额已说明，未视作等额相符",p["period"]])
    for item in outstanding(w):
        r=item["record"];put("未达明细",[r.id,r.role,r.account,r.day,format_amount(r.amount_minor),r.party,r.reference,item["age_days"],"是" if item["explained"] else "否"])
    for period,c in s.closings.items():
        for rid in c["outstanding"]:put("期间快照",[period,"当期未达",rid,""])
        for rid,amount in c["open_items"].items():put("期间快照",[period,"往来未结清",rid,format_amount(amount)])
        for rid,amount in c.get("advances",{}).items():put("期间快照",[period,"预收预付未核销",rid,format_amount(amount)])
        for rid,amount in c.get("credit_balances",{}).items():put("期间快照",[period,"红字可用余额",rid,format_amount(amount)])
        for rid in c["matches"]:put("期间快照",[period,"当期已确认关系",rid,""])
    for rid,p in s.allocations.items():put("往来核销",[rid,p["cash"],p["item"],format_amount(p["amount_minor"]),p["note"],p["period"],format_amount(allocation_net(s,rid)),p.get("advance","")])
    for item in open_items(w):
        r=item["record"];put("往来余额",[r.id,r.role,r.party,r.reference,r.day,format_amount(r.amount_minor),format_amount(item["allocated_minor"]),format_amount(item["credited_minor"]),format_amount(item["remaining_minor"]),item["age_days"],r.due or "",item["overdue_days"] if item["overdue_days"] is not None else "未提供到期日"])
    for rid,p in s.transfers.items():put("内部转账",[rid,p["outgoing"],w.records[p["outgoing"]].account,p["incoming"],w.records[p["incoming"]].account,format_amount(p["amount_minor"]),format_amount(p["fee_minor"]),p["evidence"],p["period"]])
    for rid,p in s.refunds.items():put("退款冲销",[rid,p["original"],p["cash"],format_amount(p["amount_minor"]),p["category"],p["allocation"],p["advance"],p["evidence"],p["period"]])
    for p in advance_rows(w):put("预收预付余额",[p["id"],p["kind"],p["cash"],p["party"],format_amount(p["amount_minor"]),format_amount(p["remaining_minor"]),p["evidence"],p["period"]])
    periods=[p for p in sorted(s.closings) if "aging" in s.closings[p]]
    if s.period not in periods:periods.append(s.period)
    for period in periods:
        if period in s.closings:
            statements=s.closings[period].get("bank_statements",{}).values()
        else:statements=[bank_statement(w,a.code) for a in w.accounts]
        for report in statements:
            account=report["account"]
            for label,key in (("银行期末余额","bank_minor"),("账面期末余额","ledger_minor"),("银行调节后余额","bank_adjusted_minor"),("账面调节后余额","ledger_adjusted_minor"),("调节后差额","difference_minor")):
                value=report[key];put("银行余额调节表",[period,account,label,format_amount(value) if value is not None else "未录入",report["status"]+" / "+report["evidence"]])
            for category,amount in report["totals"].items():put("银行余额调节表",[period,account,category,format_amount(amount),"仅计入已复核未达项目"])
            put("银行余额调节表",[period,account,"待复核项目数",report["unreviewed_count"],report["status"]])
            for item in report["items"]:
                r=w.records[item["record_id"]];b=batches[r.batch_id];raw=source_rows[b.id][r.source_id]
                put("余额调节明细",[period,account,r.id,r.role,r.day,format_amount(r.amount_minor),item["category"],"是" if item["reviewed"] else "否",item["evidence"],b.raw.filename,raw.source_line])
        report=aging_report(w,period)
        for g in report["groups"]:put("账龄汇总",[period,g["role"],g["party"],format_amount(g["remaining_minor"]),*[format_amount(g["buckets"][k]) for k in AGE_BUCKETS],format_amount(g["overdue_minor"]),format_amount(g["not_due_minor"]),format_amount(g["due_missing_minor"]),g["count"]])
        for item in report["details"]:
            r=w.records[item["record_id"]];b=batches[r.batch_id];raw=source_rows[b.id][r.source_id]
            put("账龄明细",[period,r.role,r.party,r.reference,r.id,r.day,item["due"],format_amount(item["allocated_minor"]),format_amount(item.get("credited_minor",0)),format_amount(item["remaining_minor"]),item["age_days"],item["bucket"],item["overdue_days"] if item["overdue_days"] is not None else "未提供",item["status"],b.raw.filename,raw.source_line])
    for rid,p in s.credits.items():put("红字抵减",[rid,p["credit"],p["item"],format_amount(p["amount_minor"]),p["category"],p["evidence"],p["period"]])
    for p in credit_rows(w):put("红字可用余额",[p["id"],p["role"],p["party"],p["reference"],format_amount(p["amount_minor"]),format_amount(p["applied_minor"]),format_amount(p["remaining_minor"])])
    duplicate_periods={p:c.get("duplicate_groups",[]) for p,c in s.closings.items()}
    if s.period not in duplicate_periods:duplicate_periods[s.period]=duplicate_groups(w,s)
    for period,groups in duplicate_periods.items():
        for g in groups:
            check_cancel(cancelled)
            for rid in g["rows"]:
                r=w.records[rid];b=batches[r.batch_id];raw=source_rows[b.id][r.source_id]
                put("疑似重复交易",[period,g["id"],g["status"],g["basis"],rid,r.role,r.account,r.day,format_amount(r.amount_minor),r.party,r.reference,b.raw.filename,raw.source_line,g["evidence"]])
    partner_reports=[r for c in s.closings.values() for r in c.get("partner_reports",{}).values()]
    if s.period not in s.closings:
        partner_reports.extend(counterparty_report(w,bid,state=s) for bid,p in s.partner_statements.items() if p["as_of"]==period_end(s.period).isoformat())
    for r in partner_reports:
        check_cancel(cancelled)
        put("双方往来对账汇总",[r["period"],r["batch"],r["party"],r["role"],r["as_of"],format_amount(r["own_total_minor"]),format_amount(r["partner_total_minor"]),format_amount(r["difference_minor"]),r["confirmed_count"],r["explained_count"],r["unreviewed_count"],r["evidence"]])
        for p in r["rows"]:put("双方往来对账明细",[r["period"],r["batch"],p["reference"],p["own"],p["partner"],format_amount(p["own_minor"]) if p["own_minor"] is not None else "缺单",format_amount(p["partner_minor"]) if p["partner_minor"] is not None else "缺单",format_amount(p["difference_minor"]),p["status"],p["evidence"]])
    export_advanced(w,s,put,cancelled)
    for b in w.batches:
        for n,r in enumerate(b.raw.rows):
            if n%250==0:check_cancel(cancelled)
            values=[(b.raw.headers[i] if i<len(b.raw.headers) else str(i+1),_encode(c)) for i,c in enumerate(r.cells)]
            # Split long JSON across rows instead of letting Excel silently truncate it.
            text=json.dumps(values,ensure_ascii=False,default=str)
            for pos in range(0,max(len(text),1),30000):
                put("原始记录",[b.id,b.role,b.account,b.raw.filename,b.raw.sha256,b.raw.sheet or "CSV",r.source_line,r.id,text[pos:pos+30000]])
    for e in w.events:
        text=json.dumps(e.data,ensure_ascii=False)
        for pos in range(0,max(len(text),1),30000):put("操作记录",[e.seq,e.kind,text[pos:pos+30000]])
    for sh in wb:
        sh.freeze_panes="A2";sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="193B50")
        for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=24
    temporary=None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent,suffix=".tmp",delete=False) as f:temporary=Path(f.name)
        wb.save(temporary);check_cancel(cancelled);os.replace(temporary,path)
    finally:
        wb.close()
        if temporary:temporary.unlink(missing_ok=True)
