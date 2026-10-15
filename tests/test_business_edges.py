import json
from dataclasses import replace
import pytest
from openpyxl import load_workbook
from local_reconcile.domain import LocalReconcileError
from local_reconcile.business.workspace import record_event,add_account
from local_reconcile.business.matching import suggest_matches
from local_reconcile.business.operations import outstanding,open_items
from local_reconcile.business.storage import save_workspace,load_workspace,export_workspace,import_legacy
from test_business import base,add,ids,sample,confirm


def test_export_keeps_pending_candidates_and_full_audit_text(tmp_path):
    w=sample();c=suggest_matches(w).candidates[0]
    w=record_event(w,"confirm_match",left=c.left,right=c.right,note="核"*32700)
    out=tmp_path/"完整.xlsx";export_workspace(w,out)
    wb=load_workbook(out)
    assert "待确认候选" in wb.sheetnames and "搜索范围" in wb.sheetnames
    parts=[row[2] for row in wb["操作记录"].iter_rows(min_row=2,values_only=True) if row[0]==w.events[-1].seq]
    assert "".join(parts)==json.dumps(w.events[-1].data,ensure_ascii=False)


def test_cross_account_and_opposite_direction_cannot_match():
    w=add_account(base(),"B","B")
    w=add(w,"bank",[["2026-10-15","10","甲","X",""]],"A")
    w=add(w,"ledger",[["2026-10-15","10","甲","X",""]],"B")
    assert not suggest_matches(w).candidates
    with pytest.raises(LocalReconcileError):record_event(w,"confirm_match",left=ids(w,"bank"),right=ids(w,"ledger"),note="不应跨账户")


def test_partial_allocation_cannot_exceed_invoice():
    w=add(base(),"bank",[["2026-10-15","100","甲","",""]])
    w=add(w,"receivable",[["2026-10-15","30","甲","I1",""]])
    with pytest.raises(LocalReconcileError):record_event(w,"allocate",cash=ids(w,"bank")[0],item=ids(w,"receivable")[0],amount_minor=3100,note="")


def test_party_mismatch_requires_allocation_evidence():
    w=add(base(),"bank",[["2026-10-15","10","甲","",""]])
    w=add(w,"receivable",[["2026-10-15","30","乙","I1",""]])
    with pytest.raises(LocalReconcileError):record_event(w,"allocate",cash=ids(w,"bank")[0],item=ids(w,"receivable")[0],amount_minor=1000,note="")


def test_exclusion_keeps_raw_and_cannot_exclude_used_row():
    w=sample();rid=ids(w,"bank")[0]
    with pytest.raises(LocalReconcileError):record_event(w,"exclude",row=rid,note="")
    hidden=record_event(w,"exclude",row=rid,note="虚构重复记录")
    assert rid in hidden.records and len(outstanding(hidden))==2
    restored=record_event(hidden,"undo_exclude",row=rid,note="复核恢复")
    assert len(outstanding(restored))==3
    w=confirm(w,suggest_matches(w).candidates[0])
    with pytest.raises(LocalReconcileError):record_event(w,"exclude",row=rid,note="已有确认")


def test_source_checksum_detects_edits(tmp_path):
    w=sample();p=tmp_path/"源.lrwork";save_workspace(w,p)
    data=json.loads(p.read_text(encoding="utf-8"));data["batches"][0]["raw"]["filename"]="被改过.csv"
    p.write_text(json.dumps(data),encoding="utf-8")
    with pytest.raises(LocalReconcileError):load_workspace(p)


def test_future_records_only_visible_after_advance():
    w=add(base(),"bank",[["2026-11-01","10","甲","X",""]])
    assert not outstanding(w)
    w=record_event(w,"close_period");w=record_event(w,"advance_period",period="2026-11")
    assert len(outstanding(w))==1 and not w.state.closings["2026-10"]["outstanding"]


def test_cancelled_save_preserves_existing_file(tmp_path):
    from concurrent.futures import CancelledError
    p=tmp_path/"旧.lrwork";p.write_text("old",encoding="utf-8")
    with pytest.raises(CancelledError):save_workspace(sample(),p,cancelled=lambda:True)
    assert p.read_text()=="old"


def test_legacy_reject_undo_and_original_history():
    from test_review import project,event
    from local_reconcile.review import apply_review
    s=project();s=apply_review(s,event(s,"reject",note="待查"));s=apply_review(s,event(s,"undo",note="核实后恢复"))
    w=import_legacy(s,"2026-10")
    assert not w.state.rejected and [e.data["legacy_sequence"] for e in w.events if "legacy_sequence" in e.data]==[1,2]


def test_independent_unique_references_do_not_scan_every_same_day_row(monkeypatch):
    import local_reconcile.business.matching as engine
    count=0;original=engine.relation_basis
    def measured(*args):
        nonlocal count
        count+=1
        return original(*args)
    monkeypatch.setattr(engine,"relation_basis",measured)
    rows=[["2026-10-15","10",f"客户{i}",f"REF{i}",""] for i in range(200)]
    w=add(base(),"bank",rows);w=add(w,"ledger",rows)
    assert len(suggest_matches(w).candidates)==200
    assert count<=2000
