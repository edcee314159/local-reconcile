import pytest
from dataclasses import replace
from datetime import date,timedelta
from test_business import base,add,ids
from conftest import raw_table,mapping
from local_reconcile.business.workspace import record_event,add_batch
from local_reconcile.business.storage import save_workspace,load_workspace,export_workspace
from local_reconcile.domain import LocalReconcileError


def fixture():
    w=add(base(),'bank',[["2026-10-10",'100','甲','B1',''],["2026-10-11",'-20','甲','B2','']])
    return add(w,'ledger',[["2026-10-12",'50','甲','L1',''],["2026-10-13",'-10','甲','L2','']])


def reviewed(w):
    names={'bank':('银行已收企业未收','银行已付企业未付'),'ledger':('企业已收银行未收','企业已付银行未付')}
    for r in w.records.values():
        w=record_event(w,'classify_unreconciled',row=r.id,category=names[r.role][r.amount_minor<0],evidence='回单核实')
    return record_event(w,'set_closing_balances',account='A',bank_minor=100000,ledger_minor=96000,evidence='银行对账单与账簿期末余额')


def test_bank_adjustment_four_signs_and_explicit_balances():
    from local_reconcile.business.reporting import bank_statement
    w=reviewed(fixture());r=bank_statement(w,'A')
    assert r['bank_adjusted_minor']==104000 and r['ledger_adjusted_minor']==104000
    assert r['difference_minor']==0 and r['status']=='调节后相符'
    assert len(r['items'])==4


def test_no_balance_and_unreviewed_never_mark_reconciled():
    from local_reconcile.business.reporting import bank_statement
    w=fixture();assert bank_statement(w,'A')['difference_minor'] is None
    w=record_event(w,'set_closing_balances',account='A',bank_minor=10000,ledger_minor=10000,evidence='余额证据')
    r=bank_statement(w,'A');assert r['status']=='存在待复核项目' and r['unreviewed_count']==4
    with pytest.raises(LocalReconcileError):record_event(w,'classify_unreconciled',row=ids(w,'bank')[0],category='企业已收银行未收',evidence='错误分类')
    with pytest.raises(LocalReconcileError):record_event(w,'set_closing_balances',account='X',bank_minor=0,ledger_minor=0,evidence='无此账户')
    with pytest.raises(LocalReconcileError):record_event(w,'set_closing_balances',account='A',bank_minor=True,ledger_minor=0,evidence='错误数值')


def test_confirmed_or_excluded_rows_do_not_adjust_and_pending_is_explicit():
    from local_reconcile.business.reporting import bank_statement
    w=add(base(),'bank',[["2026-10-10",'100','甲','SAME','']]);w=add(w,'ledger',[["2026-10-10",'100','甲','SAME','']])
    b,l=ids(w,'bank')[0],ids(w,'ledger')[0]
    w=record_event(w,'classify_unreconciled',row=b,category='银行已收企业未收',evidence='初次确认')
    w=record_event(w,'confirm_match',left=[b],right=[l],note='双方已入账')
    assert not bank_statement(w,'A')['items']
    with pytest.raises(LocalReconcileError):record_event(w,'classify_unreconciled',row=b,category='银行已收企业未收',evidence='已经匹配')


def aging_fixture():
    end=date(2026,10,31);rows=[]
    for n,age in enumerate((0,30,31,60,61,90,91)):
        rows.append([(end-timedelta(days=age)).isoformat(),'100','甲',f'I{n}','', '2026-10-30' if n%2==0 else ''])
    raw=raw_table('ledger',rows,headers=('日期','金额','对方','编号','摘要','到期日'))
    w=add_batch(base(),'receivable','',raw,mapping(declared_account='OPEN-ITEMS'),due_col=5)
    return add(w,'payable',[["2026-10-10",'80','甲','P1','']])


def test_aging_exact_boundaries_and_no_netting():
    from local_reconcile.business.reporting import aging_report
    r=aging_report(aging_fixture());assert len(r['groups'])==2
    a=next(g for g in r['groups'] if g['role']=='receivable')
    assert a['buckets']=={'0—30天':20000,'31—60天':20000,'61—90天':20000,'90天以上':10000}
    assert a['overdue_minor']==40000 and a['due_missing_minor']==30000
    assert len(r['details'])==8


def test_aging_uses_net_partial_allocation_refund_and_closed_snapshot():
    from test_settlements import cashbook,refund
    from local_reconcile.business.reporting import aging_report
    w,o,c,i=cashbook();w=record_event(w,'allocate',cash=o,item=i,amount_minor=10000,note='')
    assert aging_report(w)['details']==[]
    w=refund(w,o,c,allocation=next(iter(w.state.allocations)))
    assert aging_report(w)['details'][0]['remaining_minor']==3000
    w=record_event(w,'close_period');w=record_event(w,'advance_period',period='2026-11')
    w=record_event(w,'undo_refund',target=next(iter(w.state.refunds)),note='更正')
    assert not aging_report(w)['details'] and aging_report(w,'2026-10')['details'][0]['remaining_minor']==3000


def test_monthly_reports_roundtrip_export_and_history_immutable(tmp_path):
    from local_reconcile.business.reporting import bank_statement,aging_report
    w=reviewed(fixture());old=bank_statement(w,'A');w=record_event(w,'close_period');w=record_event(w,'advance_period',period='2026-11')
    assert bank_statement(w,'A')['difference_minor'] is None
    w=record_event(w,'set_closing_balances',account='A',bank_minor=0,ledger_minor=0,evidence='下月余额')
    assert bank_statement(w,'A','2026-10')==old
    p=tmp_path/'报表.lrwork';save_workspace(w,p);restored=load_workspace(p)
    assert bank_statement(restored,'A','2026-10')==old
    export_workspace(restored,tmp_path/'月末.xlsx')
    from openpyxl import load_workbook
    b=load_workbook(tmp_path/'月末.xlsx');assert {'银行余额调节表','余额调节明细','账龄汇总','账龄明细'}<=set(b.sheetnames);b.close()


def test_pending_classification_is_not_adjusted_and_can_be_removed():
    from local_reconcile.business.reporting import bank_statement
    w=reviewed(fixture());rid=ids(w,'bank')[0]
    w=record_event(w,'classify_unreconciled',row=rid,category='待查',evidence='仍需核实')
    assert bank_statement(w,'A')['unreviewed_count']==1
    assert bank_statement(w,'A')['totals']['银行已收企业未收']==0
    w=record_event(w,'unclassify_unreconciled',row=rid,note='撤销分类')
    assert bank_statement(w,'A')['unreviewed_count']==1
    w=record_event(w,'exclude',row=rid,note='经核实为重复来源')
    assert len(bank_statement(w,'A')['items'])==3


def test_zero_and_negative_balances_are_real_values_and_history_is_copy():
    from local_reconcile.business.reporting import bank_statement
    w=record_event(base(),'set_closing_balances',account='A',bank_minor=-100,ledger_minor=-100,evidence='透支余额')
    assert bank_statement(w,'A')['status']=='调节后相符'
    w=record_event(w,'close_period');r=bank_statement(w,'A');r['totals']['企业已收银行未收']=10000
    assert bank_statement(w,'A')['totals']['企业已收银行未收']==0
    with pytest.raises(LocalReconcileError):record_event(w,'set_closing_balances',account='A',bank_minor=0,ledger_minor=0,evidence='已结账不可改')


def test_aging_due_today_and_future_date_are_not_overdue():
    from local_reconcile.business.reporting import aging_report
    raw=raw_table('ledger',[["2026-10-10",'100','甲','DUE','',"2026-10-31"],["2026-11-01",'100','甲','FUTURE','',"2026-11-15"]],headers=('日期','金额','对方','编号','摘要','到期日'))
    w=add_batch(base(),'receivable','',raw,mapping(declared_account='OPEN-ITEMS'),due_col=5)
    r=aging_report(w);assert len(r['details'])==1
    assert r['details'][0]['overdue_days']==0 and r['groups'][0]['not_due_minor']==10000


def test_later_account_does_not_change_closed_statement_membership():
    from local_reconcile.business.workspace import add_account
    from local_reconcile.business.reporting import bank_statement
    w=reviewed(fixture());w=record_event(w,'close_period');w=record_event(w,'advance_period',period='2026-11')
    before=w.state.closings['2026-10'];w=add_account(w,'B','次月新账户')
    assert w.state.closings['2026-10']==before
    with pytest.raises(LocalReconcileError):bank_statement(w,'B','2026-10')
