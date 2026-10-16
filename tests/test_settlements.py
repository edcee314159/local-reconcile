import pytest
from test_business import base, add, ids, confirm
from local_reconcile.business.workspace import record_event
from local_reconcile.business.operations import open_items, cash_remaining
from local_reconcile.business.matching import suggest_matches
from local_reconcile.business.storage import save_workspace, load_workspace, export_workspace
from local_reconcile.domain import LocalReconcileError


def cashbook(sign=1):
    w=add(base(),'bank',[["2026-10-10",str(sign*100),'甲','PAY',''],["2026-10-12",str(-sign*30),'甲','RETURN','']])
    w=add(w,'receivable' if sign>0 else 'payable',[["2026-10-11",'100','甲','INV','']])
    return w,*ids(w,'bank'),ids(w,'receivable' if sign>0 else 'payable')[0]


def refund(w,original,cash,amount=3000,**extra):
    return record_event(w,'refund',original=original,cash=cash,amount_minor=amount,category='退款',evidence='回单已核实',**extra)


@pytest.mark.parametrize('sign',[1,-1])
def test_partial_refund_reopens_allocated_item(sign):
    w,o,c,i=cashbook(sign)
    w=record_event(w,'allocate',cash=o,item=i,amount_minor=10000,note='核实')
    a=next(iter(w.state.allocations));w=refund(w,o,c,allocation=a)
    assert open_items(w)[0]['remaining_minor']==3000
    assert cash_remaining(w,o)==0 and cash_remaining(w,c)==0
    with pytest.raises(LocalReconcileError):refund(w,o,c,allocation=a,amount=1)
    with pytest.raises(LocalReconcileError):record_event(w,'undo_allocate',target=a,note='')
    w=record_event(w,'undo_refund',target=next(iter(w.state.refunds)),note='更正')
    assert open_items(w)[0]['remaining_minor']==0
    assert cash_remaining(w,c)==3000


def test_unallocated_refund_cannot_be_spent_twice_or_reverse_date():
    w,o,c,i=cashbook();w=refund(w,o,c)
    assert cash_remaining(w,o)==7000
    with pytest.raises(LocalReconcileError):record_event(w,'allocate',cash=o,item=i,amount_minor=8000,note='')
    with pytest.raises(LocalReconcileError):refund(w,c,o,amount=100)
    with pytest.raises(LocalReconcileError):record_event(w,'exclude',row=c,note='不可排除')


@pytest.mark.parametrize('sign',[1,-1])
def test_advance_partial_settlement_and_refund(sign):
    from local_reconcile.business.settlements import advance_rows
    w,o,c,i=cashbook(sign)
    w=record_event(w,'register_advance',cash=o,amount_minor=10000,evidence='预收预付约定')
    a=next(iter(w.state.advances));assert cash_remaining(w,o)==0
    with pytest.raises(LocalReconcileError):record_event(w,'allocate',cash=o,item=i,amount_minor=1,note='')
    w=record_event(w,'allocate_advance',advance=a,item=i,amount_minor=6000,note='核实')
    assert advance_rows(w)[0]['remaining_minor']==4000
    w=refund(w,o,c,advance=a)
    assert advance_rows(w)[0]['remaining_minor']==1000
    assert open_items(w)[0]['remaining_minor']==4000
    with pytest.raises(LocalReconcileError):record_event(w,'allocate_advance',advance=a,item=i,amount_minor=2000,note='')
    with pytest.raises(LocalReconcileError):record_event(w,'undo_advance',target=a,note='')
    assert cash_remaining(w,o)==0


def test_refund_of_advance_allocation_restores_invoice_not_spendable_advance():
    from local_reconcile.business.settlements import advance_rows
    w,o,c,i=cashbook();w=record_event(w,'register_advance',cash=o,amount_minor=10000,evidence='协议')
    a=next(iter(w.state.advances));w=record_event(w,'allocate_advance',advance=a,item=i,amount_minor=10000,note='')
    w=refund(w,o,c,allocation=next(iter(w.state.allocations)))
    assert open_items(w)[0]['remaining_minor']==3000
    assert advance_rows(w)[0]['remaining_minor']==0 and cash_remaining(w,o)==0


def test_refunds_roundtrip_export_and_month_snapshot(tmp_path):
    w,o,c,i=cashbook();w=record_event(w,'allocate',cash=o,item=i,amount_minor=10000,note='')
    a=next(iter(w.state.allocations));w=record_event(w,'close_period');w=record_event(w,'advance_period',period='2026-11')
    w=refund(w,o,c,allocation=a)
    assert w.state.closings['2026-10']['open_items'][i]==0
    path=tmp_path/'退款.lrwork';save_workspace(w,path);r=load_workspace(path)
    assert r.state.refunds==w.state.refunds and open_items(r)[0]['remaining_minor']==3000
    export_workspace(r,tmp_path/'报告.xlsx')
    from openpyxl import load_workbook
    book=load_workbook(tmp_path/'报告.xlsx');assert {'退款冲销','预收预付余额'}<=set(book.sheetnames);book.close()


def test_many_to_many_suggestion_and_confirmation():
    w=add(base(),'bank',[["2026-10-15",str(a),'甲','GROUP',''] for a in (60,40)])
    w=add(w,'ledger',[["2026-10-15",str(a),'甲','GROUP',''] for a in (70,30)])
    choices=[c for c in suggest_matches(w).candidates if len(c.left)==len(c.right)==2]
    assert len(choices)==1
    w=confirm(w,choices[0]);assert len(w.state.matches)==1
    with pytest.raises(LocalReconcileError):confirm(w,choices[0])


def test_many_to_many_needs_shared_basis_and_obeys_budget():
    from dataclasses import replace
    w=add(base(),'bank',[["2026-10-15",str(a),'甲','X',''] for a in (60,40)])
    w=add(w,'ledger',[["2026-10-15",str(a),'甲','Y',''] for a in (70,30)])
    assert not suggest_matches(w).candidates
    with pytest.raises(LocalReconcileError):record_event(w,'confirm_match',left=ids(w,'bank'),right=ids(w,'ledger'),note='')
    w=add(base(),'bank',[["2026-10-15",str(a),'甲','GROUP',''] for a in (60,40)])
    w=add(w,'ledger',[["2026-10-15",str(a),'甲','GROUP',''] for a in (70,30)])
    w=replace(w,settings=replace(w.settings,project_budget=1))
    assert suggest_matches(w).ranges


def test_undo_refund_after_reallocation_cannot_overallocate():
    w,o,c,i=cashbook();w=add(w,'bank',[["2026-10-15",'30','甲','REPAY','']])
    other=ids(w,'bank')[-1];w=record_event(w,'allocate',cash=o,item=i,amount_minor=10000,note='')
    w=refund(w,o,c,allocation=next(iter(w.state.allocations)))
    w=record_event(w,'allocate',cash=other,item=i,amount_minor=3000,note='')
    with pytest.raises(LocalReconcileError):record_event(w,'undo_refund',target=next(iter(w.state.refunds)),note='恢复会超额')


def test_prepaid_must_match_party_direction_and_later_invoice():
    w,o,c,i=cashbook();w=add(w,'receivable',[["2026-10-09",'100','甲','OLD',''],["2026-10-15",'100','乙','WRONG','']])
    w=record_event(w,'register_advance',cash=o,amount_minor=10000,evidence='预收协议')
    a=next(iter(w.state.advances))
    for item in ids(w,'receivable')[1:]:
        with pytest.raises(LocalReconcileError):record_event(w,'allocate_advance',advance=a,item=item,amount_minor=1000,note='')
    with pytest.raises(LocalReconcileError):record_event(w,'register_advance',cash=o,amount_minor=1,evidence='重复')
    with pytest.raises(LocalReconcileError):record_event(w,'exclude',row=o,note='已占用')


def test_advance_save_replay_export_and_snapshot(tmp_path):
    from local_reconcile.business.settlements import advance_rows
    w,o,c,i=cashbook(-1);w=record_event(w,'register_advance',cash=o,amount_minor=10000,evidence='预付款协议')
    a=next(iter(w.state.advances));w=record_event(w,'allocate_advance',advance=a,item=i,amount_minor=6000,note='')
    w=record_event(w,'close_period');w=record_event(w,'advance_period',period='2026-11');w=refund(w,o,c,advance=a)
    assert w.state.closings['2026-10']['advances'][a]==4000
    p=tmp_path/'预付款.lrwork';save_workspace(w,p);restored=load_workspace(p)
    assert advance_rows(restored)[0]['remaining_minor']==1000
    export_workspace(restored,tmp_path/'预付款.xlsx')
    from openpyxl import load_workbook
    b=load_workbook(tmp_path/'预付款.xlsx');assert b['预收预付余额'].cell(2,6).value=='10.00';b.close()


@pytest.mark.parametrize('category',['冲销','退票'])
def test_full_reversal_and_returned_cheque(category):
    w=add(base(),'bank',[["2026-10-01",'100','甲','ORIG',''],["2026-10-02",'-100','甲','RETURN','']]);o,c=ids(w,'bank')
    w=record_event(w,'refund',original=o,cash=c,amount_minor=10000,category=category,evidence='已核实银行通知')
    assert cash_remaining(w,o)==cash_remaining(w,c)==0
    with pytest.raises(LocalReconcileError):refund(w,o,c,amount=1)


def test_many_to_many_negative_party_basis_and_cent_precision():
    w=add(base(),'bank',[["2026-10-15",a,'供应商','', ''] for a in ('-0.06','-0.04')])
    w=add(w,'ledger',[["2026-10-15",a,'供应商','', ''] for a in ('-0.07','-0.03')])
    c=next(c for c in suggest_matches(w).candidates if len(c.left)==len(c.right)==2)
    assert c.amount_minor==-10 and c.basis=='同一交易对方'
    assert confirm(w,c).state.matches


def test_many_to_many_pool_overflow_and_cancel():
    from concurrent.futures import CancelledError
    w=add(base(),'bank',[["2026-10-15",'60','甲','POOL',''],["2026-10-15",'40','甲','POOL','']])
    w=add(w,'ledger',[["2026-10-15",'6','甲','POOL',''] for _ in range(17)])
    a=suggest_matches(w);assert a.ranges and not a.candidates
    assert set(ids(w,'bank')+ids(w,'ledger'))<=set().union(*(set(r.rows) for r in a.ranges))
    with pytest.raises(CancelledError):suggest_matches(w,cancelled=lambda:True)


def test_many_to_many_shared_party_with_missing_references():
    w=add(base(),'bank',[["2026-10-15",str(a),'甲','X',''] for a in (60,40)])
    w=add(w,'ledger',[["2026-10-15",str(a),'甲','',''] for a in (70,30)])
    c=[c for c in suggest_matches(w).candidates if len(c.left)==len(c.right)==2]
    assert len(c)==1 and c[0].basis=='同一交易对方'
