from PySide6.QtCore import Qt,QDate
from test_ui_flow import app
from local_reconcile.ui.business_window import BusinessWindow
from local_reconcile.business.advanced_demo import advanced_demo,run_advanced_smoke
from local_reconcile.business.net_settlements import net_settlement_reports
from local_reconcile.business.cutoff import cutoff_report
from openpyxl import load_workbook


def test_real_net_and_cutoff_ui_actions_history_and_sources(app):
    w=BusinessWindow()
    try:
        w.set_workspace(advanced_demo());assert w.tabs.count()==17
        assert w.net_bank_list.count()==1
        w.net_bank_list.item(0).setSelected(True);w.net_evidence.setText('原订单与退款及费用已核对')
        w.confirm_net_selected();assert net_settlement_reports(w.workspace)[0]['review_valid']
        row=next(p for p in w.cutoff_rows if p['reference']=='ACCEPT-1')
        w.cutoff_business.setCurrentIndex(w.cutoff_business.findData(row['business']))
        w.cutoff_explanation.setText('已核查尚缺30元，本期待补核算');w.explain_cutoff_selected()
        row=next(p for p in cutoff_report(w.workspace)['rows'] if p['reference']=='ACCEPT-1')
        assert row['delta_minor']==-3000 and row['explanation_valid']
        pending=next(p for p in w.cutoff_rows if p['reference']=='UNREVIEWED-1')
        w.cutoff_business.setCurrentIndex(w.cutoff_business.findData(pending['business']))
        w.cutoff_evidence.setText('核实业务按发生日归属');w.save_cutoff_rule_selected()
        assert w.workspace.state.cutoff_rules[pending['business']]['method']=='once'
        w.apply('close_period');w.apply('advance_period',period='2026-11')
        w.net_period.setCurrentIndex(w.net_period.findData('2026-10'))
        w.cutoff_period.setCurrentIndex(w.cutoff_period.findData('2026-10'))
        assert not w.net_confirm_button.isEnabled() and not w.cutoff_save_button.isEnabled()
        assert w.net_source_table.model().rowCount()>=5 and w.cutoff_table.model().rowCount()==4
    finally:w.dirty=False;w.close()


def test_advanced_export_and_frozen_smoke_values(tmp_path):
    r=run_advanced_smoke(tmp_path)
    assert r['net_settlement_minor']==9200 and r['net_settlement_confirmed']
    assert r['cutoff_missing_minor']==3000 and r['cutoff_monthly_minor']==10000 and r['cutoff_unmatched_postings']==1
    assert r['advanced_history_preserved']
    wb=load_workbook(tmp_path/'advanced-report.xlsx',data_only=True)
    assert {'净到账还原汇总','结算组成明细','结算订单核验','结算到账流水','月末核算检查','核算记账依据'}<=set(wb.sheetnames)
    rows=list(wb['月末核算检查'].values)
    assert any('差异已说明仍待复核' in p for p in rows)
    assert any('未确认' in p for p in rows)
    assert any('记账记录未关联业务' in p for p in rows)
    wb.close()


def test_zero_net_confirmation_and_undo_are_accessible_in_ui(app):
    from test_business import base,add
    from test_net_settlements import lines
    data=add(base(),'settlement_order',[["2026-10-01",'100','平台甲','O1','']])
    data=lines(data,[['2026-10-15','100','平台甲','S1','','B1','销售','O1'],['2026-10-16','-100','平台甲','R1','','B1','退款','O1']])
    w=BusinessWindow()
    try:
        w.set_workspace(data);assert w.net_bank_list.count()==0
        w.net_evidence.setText('全额退回，无实际资金结算');w.confirm_net_selected()
        assert w.net_undo_button.isEnabled()
        w.undo_net_selected();assert not net_settlement_reports(w.workspace)[0]['review_valid']
    finally:w.dirty=False;w.close()
