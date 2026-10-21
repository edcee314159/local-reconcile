from test_ui_flow import app
from local_reconcile.ui.business_window import BusinessWindow
from local_reconcile.business.commercial_demo import commercial_demo, run_commercial_smoke
from local_reconcile.business.operations import open_items
from local_reconcile.business.storage import export_workspace, save_workspace, load_workspace
from local_reconcile.business.counterparty import counterparty_report
from openpyxl import load_workbook


def test_commercial_ui_credit_duplicate_and_counterparty(app, tmp_path):
    w=BusinessWindow()
    try:
        w.set_workspace(commercial_demo())
        credit=next(r for r in w.workspace.records.values() if r.role=='receivable_credit')
        w.credit_source.setCurrentIndex(w.credit_source.findData(credit.id))
        w.credit_amount.setText('20');w.credit_evidence.setText('退货凭据')
        w.apply_credit_selected()
        assert next(p['remaining_minor'] for p in open_items(w.workspace) if p['record'].reference=='INV-C1')==8000
        assert w.credit_source_table.model().rowCount()==2
        w.duplicate_table.selectRow(0);w.render_duplicate_members()
        assert w.duplicate_member_table.model().rowCount()==2
        w.duplicate_member_table.selectRow(1);w.duplicate_evidence.setText('重复导出回单')
        w.exclude_duplicate_selected();assert not w.duplicate_groups
        bid=next(b.id for b in w.workspace.batches if b.role=='partner_receivable')
        w.partner_batch.setCurrentIndex(w.partner_batch.findData(bid))
        w.partner_as_of.setText('2026-10-31');w.partner_statement_evidence.setText('月末签收')
        w.register_partner_selected()
        ix=next(i for i,p in enumerate(w.partner_rows) if p['reference']=='INV-C1')
        w.partner_table.selectRow(ix);w.partner_decision_evidence.setText('核对无误');w.decide_partner_selected('confirmed')
        assert counterparty_report(w.workspace,bid)['confirmed_count']==1
        w.apply('close_period');w.apply('advance_period',period='2026-11')
        w.partner_period.setCurrentIndex(w.partner_period.findData('2026-10'))
        w.partner_batch.setCurrentIndex(w.partner_batch.findData(bid))
        assert not w.partner_confirm_button.isEnabled() and w.partner_table.model().rowCount()==3
        assert '2026-10' in w.partner_status.text()
        path=tmp_path/'商业.lrwork';save_workspace(w.workspace,path)
        assert load_workspace(path).state==w.workspace.state
    finally:w.dirty=False;w.close()


def test_commercial_export_values_history_and_missing_sides(tmp_path):
    result=run_commercial_smoke(tmp_path)
    assert result['credit_receivable_remaining_minor']==8000 and result['credit_payable_remaining_minor']==15000
    assert result['duplicate_use_blocked'] and result['duplicate_groups_after_exclusion']==0
    assert result['partner_confirmed_rows']==1 and result['partner_explained_rows']==3 and result['partner_history_preserved']
    wb=load_workbook(tmp_path/'commercial-report.xlsx',data_only=True)
    assert {'红字抵减','红字可用余额','疑似重复交易','双方往来对账汇总','双方往来对账明细'}<=set(wb.sheetnames)
    values=list(wb['双方往来对账明细'].values)
    assert len(values)==5
    assert any('缺单' in row for row in values[1:])
    assert any('差异已说明仍未核平' in row for row in values[1:])
    header=list(wb['往来余额'].values)[0]
    assert '资金净核销（元）' in header and '红字抵减（元）' in header
    wb.close()
