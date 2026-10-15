from dataclasses import replace
from datetime import date
import pytest
from conftest import raw_table, mapping
from local_reconcile.domain import LocalReconcileError
from local_reconcile.business.workspace import new_workspace, add_account, add_batch, record_event
from local_reconcile.business.matching import suggest_matches
from local_reconcile.business.operations import open_items, outstanding, suggest_transfers
from local_reconcile.business.storage import save_workspace, load_workspace, export_workspace, import_legacy


def base():
    return add_account(new_workspace("2026-10"), "A", "银行A")


def add(w, role, rows, account="A"):
    raw = raw_table("bank" if role == "bank" else "ledger", rows)
    return add_batch(w, role, account if role in ("bank","ledger") else "", raw,
                     mapping(declared_account=account if role in ("bank","ledger") else "OPEN-ITEMS"))


def ids(w, role):
    return [r.id for r in w.records.values() if r.role == role]


def sample():
    w = add(base(),"bank",[["2026-10-15","10000","甲","X",""]])
    return add(w,"ledger",[["2026-10-15","6000","甲","X",""],["2026-10-15","4000","甲","X",""]])


def confirm(w, c):
    return record_event(w,"confirm_match",left=list(c.left),right=list(c.right),note="已核实")


def test_combo_confirm_undo_occupancy():
    w=sample(); a=suggest_matches(w)
    assert len(a.candidates)==1
    c=a.candidates[0]; assert len(c.left)==1 and len(c.right)==2 and c.amount_minor==1000000
    w=confirm(w,c)
    with pytest.raises(LocalReconcileError): confirm(w,c)
    w=record_event(w,"undo_match",target=next(iter(w.state.matches)),note="更正")
    assert not w.state.matches and len(outstanding(w))==3


def test_reverse_combo_and_exact_cents():
    w=add(base(),"bank",[["2026-10-15","0.10","甲","",""],["2026-10-15","0.20","甲","",""]])
    w=add(w,"ledger",[["2026-10-15","0.30","甲","",""]])
    c=suggest_matches(w).candidates[0]
    assert len(c.left)==2 and c.amount_minor==30
    w2=add(base(),"ledger",[["2026-10-15","0.31","甲","",""]])
    w2=add(w2,"bank",[["2026-10-15","0.10","甲","",""],["2026-10-15","0.20","甲","",""]])
    assert not suggest_matches(w2).candidates


def test_conflicting_references_no_fallback():
    w=add(base(),"bank",[["2026-10-15","10","甲","X",""]])
    w=add(w,"ledger",[["2026-10-15","6","甲","Y",""],["2026-10-15","4","甲","Z",""]])
    assert not suggest_matches(w).candidates


def test_missing_basis_and_opposite_direction_do_not_combine():
    w=add(base(),"bank",[["2026-10-15","10","","",""]])
    w=add(w,"ledger",[["2026-10-15","6","","",""],["2026-10-15","4","","",""]])
    assert not suggest_matches(w).candidates
    with pytest.raises(LocalReconcileError):
        record_event(w,"confirm_match",left=ids(w,"bank"),right=ids(w,"ledger"),note="")
    assert confirm(w,type("C",(),{"left":ids(w,"bank"),"right":ids(w,"ledger")})()).state.matches


def test_ambiguity_preserved_and_no_auto_confirmation():
    w=add(base(),"bank",[["2026-10-15","10","甲","",""]])
    w=add(w,"ledger",[["2026-10-15",str(n),"甲","",""] for n in (6,4,7,3)])
    assert len(suggest_matches(w).candidates)==2 and not w.state.matches


def test_search_limits_keep_all_source_ids():
    w=add(base(),"bank",[["2026-10-15","10","甲","",""]])
    w=add(w,"ledger",[["2026-10-15","1","甲","",""]]*17)
    a=suggest_matches(w)
    assert a.ranges and set(ids(w,"ledger")) <= set().union(*(set(r.rows) for r in a.ranges))


def test_group_cap_applies_to_whole_component():
    w=add(base(),"bank",[["2026-10-15","10",f"甲{i}","X",""] for i in range(15)])
    w=add(w,"ledger",[["2026-10-15","10","乙","X",""] for _ in range(10)])
    a=suggest_matches(w)
    assert not a.candidates and any(len(r.rows)==25 for r in a.ranges)


def test_cancel():
    from concurrent.futures import CancelledError
    with pytest.raises(CancelledError):suggest_matches(sample(),cancelled=lambda:True)


def test_import_duplicate_and_mapping_validation():
    w=base(); raw=raw_table("bank",[["2026-10-15","10","甲","X",""]]); m=mapping(declared_account="A")
    w=add_batch(w,"bank","A",raw,m)
    with pytest.raises(LocalReconcileError):add_batch(w,"bank","A",raw,m)
    with pytest.raises(LocalReconcileError):add_batch(w,"bank","A",raw,replace(m,confirmed=False))
    with pytest.raises(LocalReconcileError):add_account(w,"A","重复")


def test_explain_difference_is_not_equal_match():
    w=add(base(),"bank",[["2026-10-15","9980","甲","X",""]])
    w=add(w,"ledger",[["2026-10-15","10000","甲","X",""]])
    b,l=ids(w,"bank")[0],ids(w,"ledger")[0]
    with pytest.raises(LocalReconcileError):record_event(w,"explain",bank=b,ledger=l,category="手续费",evidence="")
    w=record_event(w,"explain",bank=b,ledger=l,category="手续费",evidence="虚构结算单 F001")
    assert next(iter(w.state.explanations.values()))["delta_minor"]==-2000
    assert not w.state.matches and len(outstanding(w))==2
    with pytest.raises(LocalReconcileError):record_event(w,"confirm_match",left=[b],right=[l],note="差额已说明")


def test_partial_allocation_and_undo():
    w=add(base(),"bank",[["2026-10-15","3000","甲","X",""]])
    w=add(w,"receivable",[["2026-10-01","10000","甲","INV-1",""]])
    b,i=ids(w,"bank")[0],ids(w,"receivable")[0]
    w=record_event(w,"allocate",cash=b,item=i,amount_minor=300000,note="核实回款")
    assert open_items(w)[0]["remaining_minor"]==700000
    with pytest.raises(LocalReconcileError):record_event(w,"allocate",cash=b,item=i,amount_minor=1,note="超额")
    w=record_event(w,"undo_allocate",target=next(iter(w.state.allocations)),note="撤销")
    assert open_items(w)[0]["remaining_minor"]==1000000


def test_payable_direction_and_party_guard():
    w=add(base(),"bank",[["2026-10-15","-30","甲","",""],["2026-10-15","30","乙","",""]])
    w=add(w,"payable",[["2026-10-01","100","甲","P1",""]])
    i=ids(w,"payable")[0]; b=ids(w,"bank")
    with pytest.raises(LocalReconcileError):record_event(w,"allocate",cash=b[1],item=i,amount_minor=3000,note="已核实")
    w=record_event(w,"allocate",cash=b[0],item=i,amount_minor=3000,note="")
    assert open_items(w)[0]["remaining_minor"]==7000


def test_month_close_is_historical_snapshot():
    w=sample(); w=record_event(w,"close_period")
    before=w.state.closings["2026-10"]
    with pytest.raises(LocalReconcileError):confirm(w,suggest_matches(w).candidates[0])
    w=record_event(w,"advance_period",period="2026-11")
    w=confirm(w,suggest_matches(w).candidates[0])
    assert not outstanding(w)
    assert w.state.closings["2026-10"]==before and len(before["outstanding"])==3


def test_backdated_late_import_requires_correction_workspace():
    w=record_event(base(),"close_period");w=record_event(w,"advance_period",period="2026-11")
    with pytest.raises(LocalReconcileError):add(w,"bank",[["2026-10-20","10","甲","NEW",""]])


def test_multi_account_transfer_and_independent_bank_match():
    w=add_account(base(),"B","银行B")
    w=add(w,"bank",[["2026-10-15","-100","内部","T1",""]],"A")
    w=add(w,"bank",[["2026-10-16","100","内部","T1",""]],"B")
    w=add(w,"ledger",[["2026-10-15","-100","内部","T1",""]],"A")
    assert len(suggest_transfers(w))==1
    out,inc=[r.id for r in w.records.values() if r.role=="bank"]
    w=record_event(w,"confirm_transfer",outgoing=out,incoming=inc,fee_minor=0,evidence="本公司账户")
    w=confirm(w,suggest_matches(w).candidates[0])
    assert len(w.state.transfers)==1 and len(w.state.matches)==1
    assert len(outstanding(w))==1


def test_transfer_fee_and_no_allocation_of_internal_cash():
    w=add_account(base(),"B","银行B")
    w=add(w,"bank",[["2026-10-15","-100.20","内部","T1",""]],"A")
    w=add(w,"bank",[["2026-10-15","100","内部","T1",""]],"B")
    w=add(w,"receivable",[["2026-10-15","100","内部","INV1",""]])
    o,i=ids(w,"bank")
    with pytest.raises(LocalReconcileError):record_event(w,"confirm_transfer",outgoing=o,incoming=i,fee_minor=20,evidence="")
    w=record_event(w,"confirm_transfer",outgoing=o,incoming=i,fee_minor=20,evidence="虚构手续费回单")
    with pytest.raises(LocalReconcileError):record_event(w,"allocate",cash=i,item=ids(w,"receivable")[0],amount_minor=10000,note="不应核销")


def test_storage_replay_report_and_tamper(tmp_path):
    from openpyxl import load_workbook
    import json
    w=sample();w=confirm(w,suggest_matches(w).candidates[0])
    p=tmp_path/"业务.lrwork";save_workspace(w,p);restored=load_workspace(p)
    assert restored.state==w.state and restored.records==w.records
    report=tmp_path/"业务.xlsx";export_workspace(w,report)
    wb=load_workbook(report,data_only=True)
    assert {"已确认对账","对账组成明细","差额说明","未达明细","往来余额","内部转账","原始记录"}<=set(wb.sheetnames)
    assert wb["对账组成明细"].max_row==4
    content=json.loads(p.read_text(encoding="utf-8"))
    content["events"][-1]["seq"]=99;p.write_text(json.dumps(content),encoding="utf-8")
    with pytest.raises(LocalReconcileError):load_workspace(p)


def test_legacy_import_preserves_review():
    from test_review import project,event
    from local_reconcile.review import apply_review
    s=apply_review(project(),event(project(),"confirm"))
    w=import_legacy(s,"2026-10")
    assert len(w.state.matches)==1 and not outstanding(w)
