from concurrent.futures import CancelledError
from pathlib import Path
import pytest
from openpyxl import load_workbook
from test_review import project,event
from local_reconcile.review import apply_review
from local_reconcile.domain import LocalReconcileError
from local_reconcile.export_xlsx import export_report
from local_reconcile.project_store import save_project

def test_report_totals_and_provenance(tmp_path):
    s=project();s=apply_review(s,event(s,"confirm"))
    p=tmp_path/"核对 报告.xlsx";export_report(s,p)
    w=load_workbook(p,data_only=True)
    assert len(w.sheetnames)==8
    assert w["核对总览"]["B9"].value==10
    rows=list(w["已确认关系"].values)
    assert "bank.csv" in rows[1] and 2 in rows[1]
    assert w["银行原始明细"].max_row==2
    w.close()

def test_untrusted_cells_are_text(tmp_path):
    s=project(b=[["2026-10-15","10","甲","A","=SUM(1,2)"]],l=[["2026-10-15","10","甲","A","@SUM(A1)"]])
    p=tmp_path/"safe.xlsx";export_report(s,p)
    w=load_workbook(p,data_only=False)
    values=[c for row in w["银行原始明细"] for c in row if c.value=="=SUM(1,2)"]
    assert len(values)==1 and values[0].data_type=="s"
    assert all(c.data_type!="f" for sheet in w for row in sheet for c in row)
    w.close()

def test_export_atomic_failure(tmp_path,monkeypatch):
    from local_reconcile import export_xlsx
    p=tmp_path/"a.xlsx";export_report(project(),p);before=p.read_bytes()
    def fail(*args):raise PermissionError()
    monkeypatch.setattr(export_xlsx.os,"replace",fail)
    with pytest.raises(LocalReconcileError):
        export_report(project(),p)
    assert p.read_bytes()==before
    assert not list(tmp_path.glob("*.tmp"))

def test_confirmed_not_candidate_total(tmp_path):
    p=tmp_path/"a.xlsx";export_report(project(),p)
    w=load_workbook(p,data_only=True)
    assert w["核对总览"]["B9"].value==0
    assert w["核对总览"]["B4"].value==1
    w.close()

def test_cancel_no_partial_report(tmp_path):
    p=tmp_path/"cancel.xlsx"
    with pytest.raises(CancelledError):
        export_report(project(),p,cancelled=lambda:True)
    assert not p.exists()

def test_never_overwrite_input(tmp_path):
    p=tmp_path/"bank.csv";p.write_text("original",encoding="utf-8");before=p.read_bytes()
    for save in (save_project,export_report):
        with pytest.raises(LocalReconcileError):
            save(project(),p)
        assert p.read_bytes()==before
