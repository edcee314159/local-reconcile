"""Packaged audit demo, immutable historical export and safe spreadsheet delivery."""
from collections import defaultdict
from concurrent.futures import CancelledError
from copy import deepcopy
import json
import pytest
from local_reconcile.business.workspace import record_event
from local_reconcile.domain import never_cancel


def demo():
    from local_reconcile.business.audit_demo import audit_demo
    return audit_demo()


def export_rows(w):
    from local_reconcile.business.audit_export import HEADERS, export_audit
    rows = defaultdict(list)
    export_audit(w, w.state, lambda name, row: rows[name].append(row), never_cancel)
    assert set(rows) <= set(HEADERS)
    assert all(len(row) == len(HEADERS[name]) for name, values in rows.items() for row in values)
    return rows


def november(w):
    return record_event(record_event(w, "close_period"), "advance_period", period="2026-11")


def test_demo_contains_confirmed_and_missing_evidence_with_three_real_stages():
    from local_reconcile.business.evidence_chain import evidence_report
    from local_reconcile.business.contract_terms import contract_report
    from local_reconcile.business.settlements import item_used
    w = demo(); groups = evidence_report(w)["groups"]; contracts = contract_report(w)
    assert len(groups) == 2 and sum(p["review_valid"] for p in groups) == 1
    assert any(p["missing"] and not p["review_valid"] for p in groups)
    assert len(contracts["stages"]) == 3 and sum(p["paid_minor"] for p in contracts["stages"]) == 30000
    complete = next(p for p in contracts["contracts"] if p["stage_count"])
    assert complete["amount_minor"] == complete["planned_minor"] == 100000
    assert sorted(p["amount_minor"] for p in contracts["stages"]) == [10000, 30000, 60000]
    wrong = next(r for r in w.records.values() if r.reference == "AR-WRONG")
    assert item_used(w.state, wrong.id) == 30000 and not w.state.corrections


def test_export_rows_are_balanced_and_contain_source_and_proof_identifiers():
    w = demo(); rows = export_rows(w)
    assert rows["业务证据链"] and rows["业务凭证明细"] and rows["凭证资金关联"]
    assert len(rows["合同阶段明细"]) == 3 and len(rows["阶段收付证明"]) == 1
    proof = next(iter(w.state.stage_cash_links.values()))
    assert any(proof["cash"] in row for row in rows["阶段收付证明"])
    assert any("300.00" in row for row in rows["凭证资金关联"])
    assert any("应收" in row for row in rows["合同阶段汇总"])


def test_export_history_retains_later_undone_evidence_and_stage_cash_details():
    w = demo(); before = deepcopy(export_rows(w))
    w = november(w)
    w = record_event(w, "undo_stage_cash", target=next(iter(w.state.stage_cash_links)), note="当前期更正阶段证明")
    w = record_event(w, "undo_evidence_cash", target=next(iter(w.state.evidence_cash_links)), note="当前期更正资金证明")
    rows = export_rows(w)
    for name in ("凭证资金关联", "阶段收付证明"):
        assert rows[name] == before[name]
        assert all(row[0] == "2026-10" for row in rows[name])
    assert {row[0] for row in rows["合同阶段明细"]} == {"2026-10", "2026-11"}


def test_correction_audit_json_chunks_preserve_all_operations_and_reason():
    from local_reconcile.business.corrections import preview_correction, apply_correction
    w = demo(); allocation = next(iter(w.state.allocations)); original = w.state.allocations[allocation]
    correct = next(r.id for r in w.records.values() if r.reference == "AR-CORRECT")
    operations = [{"kind": "undo_allocate", "data": {"target": allocation, "note": "史" * 32767}},
                  {"kind": "allocate", "data": {"cash": original["cash"], "item": correct, "amount_minor": 30000, "note": "更正归属"}}]
    reason = "原核销单据选择有误"
    w = apply_correction(w, preview_correction(w, operations, reason), reason)
    rows = export_rows(w)["业务更正审计"]
    assert len(rows) >= 2 and all(len(row[-1]) <= 30000 for row in rows)
    assert json.loads("".join(row[-1] for row in rows)) == operations
    assert all(reason in row for row in rows)


def test_export_cancellation_stops_before_any_row_is_written():
    from local_reconcile.business.audit_export import export_audit
    w = demo(); writes = []
    with pytest.raises(CancelledError): export_audit(w, w.state, lambda *row: writes.append(row), lambda: True)
    assert writes == []


def test_xlsx_audit_sheets_and_formula_injection_are_safe(tmp_path):
    from openpyxl import load_workbook
    from local_reconcile.business.audit_export import HEADERS
    from local_reconcile.business.evidence_chain import evidence_report
    from local_reconcile.business.storage import export_workspace
    w = demo(); group = next(p for p in evidence_report(w)["groups"] if p["review_valid"])
    w = record_event(w, "register_evidence_rule", group=group["id"], required=list(group["required"]), contract=group["contract"], evidence="=1+1")
    w = record_event(w, "confirm_evidence_chain", group=group["id"], evidence="=2+2")
    path = tmp_path / "audit-safe.xlsx"; export_workspace(w, path)
    wb = load_workbook(path, data_only=False)
    assert set(HEADERS) <= set(wb.sheetnames)
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    values = [cell.value for row in wb["业务证据链"] for cell in row]
    assert "=1+1" in values and "=2+2" in values
    assert all(cell.data_type == "s" for row in wb["业务证据链"] for cell in row if cell.value in ("=1+1", "=2+2"))
    wb.close()


def test_smoke_delivers_history_replay_and_atomic_correction(tmp_path):
    from local_reconcile.business.audit_demo import run_audit_smoke
    from local_reconcile.business.storage import load_workspace
    from local_reconcile.business.settlements import item_used
    result = run_audit_smoke(tmp_path)
    assert result == {"evidence_chain_confirmed": True, "contract_stage_count": 3, "contract_paid_minor": 30000,
                      "correction_audit_count": 1, "audit_history_preserved": True}
    assert (tmp_path / "audit-report.xlsx").is_file()
    w = load_workspace(tmp_path / "audit-demo.lrwork")
    wrong = next(r.id for r in w.records.values() if r.reference == "AR-WRONG")
    correct = next(r.id for r in w.records.values() if r.reference == "AR-CORRECT")
    assert item_used(w.state, wrong) == 0 and item_used(w.state, correct) == 30000
    assert w.state.closings["2026-10"]["open_items"][wrong] == 70000
    assert w.state.closings["2026-10"]["open_items"][correct] == 100000
